"""Citation contracts, native hybrid results and history snapshots without services."""

from concurrent.futures import ThreadPoolExecutor
from types import SimpleNamespace
from uuid import uuid4

from langchain_core.messages import AIMessage, AIMessageChunk
import numpy as np
import pytest

from BackEnd.app.api.document import stream_retrieval_events
from BackEnd.app.chatbot.agents import Agents
from BackEnd.app.chatbot.tools import ToolList
from BackEnd.app.database.qdrant_manager import QDrant
from BackEnd.app.database.sql_manager import Supabase_Manager
from BackEnd.app.database.sql_models import Chunk, Document
from BackEnd.app.doc_extractor import extractor
from BackEnd.app.pipeline import Pipeline
from BackEnd.app.retrieval_models import (
    AppContext, MAX_CONTEXT_CHARACTERS, RetrievalHit, SourceRegistry,
    format_retrieval_context, validate_citations,
)


def hit(point_id: str = "point-1", **changes) -> RetrievalHit:
    values = dict(qdrant_point_id=point_id, chunk_id="chunk-1", document_id="doc-1",
                  file_name="notes.pdf", page=3, chunk_index=0,
                  content="Document evidence", score=0.75, retrieval_rank=1)
    return RetrievalHit(**(values | changes))


def test_source_registry_deduplicates_and_limits_context() -> None:
    registry = SourceRegistry()
    source = registry.register([hit(), hit()])
    assert source[0].citation_id == source[1].citation_id == 1
    second = registry.register([hit("point-2", content="x" * MAX_CONTEXT_CHARACTERS)])
    assert len(second[0].content) == MAX_CONTEXT_CHARACTERS - len(source[0].content)
    assert registry.register([hit("point-3")]) == []
    assert len(registry.snapshot()) == 2
    assert "SOURCE 1" in format_retrieval_context(source)


def test_registries_are_request_local_and_parallel_safe() -> None:
    registry = SourceRegistry()
    with ThreadPoolExecutor(max_workers=4) as pool:
        list(pool.map(lambda _: registry.register([hit()]), range(20)))
    assert len(registry.snapshot()) == 1
    assert SourceRegistry().register([hit()])[0].citation_id == 1


def test_validate_citations_ignores_code_and_rejects_unknown_ids() -> None:
    sources = SourceRegistry().register([hit()])
    result = validate_citations(
        "Fact [[source:1]] [[source:99]] [[source:1]]\n"
        "`[[source:2]]`\n```text\n[[source:3]]\n```",
        [source.model_dump() for source in sources],
    )
    assert result == {"cited_source_ids": [1], "invalid_source_ids": [99]}


class StubBm25:
    def embed(self, texts):
        return [SimpleNamespace(indices=np.array([1]), values=np.array([0.4])) for _ in texts]

    def query_embed(self, query):
        yield SimpleNamespace(indices=np.array([1]), values=np.array([0.4]))


class StubQdrantClient:
    def __init__(self):
        self.calls = []

    def upsert(self, **kwargs):
        self.calls.append(kwargs)

    def query_points(self, **kwargs):
        self.calls.append(kwargs)
        return SimpleNamespace(points=[
            SimpleNamespace(id="legacy-point", score=0.7, payload={"content": "old", "document_id": "doc-1"}),
            SimpleNamespace(id="new-point", score=0.5, payload={
                "content": "new", "document_id": "doc-1", "chunk_id": "sql-chunk",
                "file_name": "notes.pdf", "page": 4, "metadata_version": 2,
            }),
        ])


def qdrant() -> QDrant:
    instance = QDrant.__new__(QDrant)
    instance.client = StubQdrantClient()
    instance.bm25_model = StubBm25()
    return instance


def test_native_rrf_returns_ids_scores_and_legacy_unknown_locations() -> None:
    instance = qdrant()
    results = instance.search_hits("user-1", "question", np.array([0.1, 0.2]), "chat-1", doc_id="doc-1")
    assert results[0].chunk_id is None
    assert results[0].qdrant_point_id == "legacy-point"
    assert results[0].page is None
    assert results[1].chunk_id == "sql-chunk"
    assert [result.score for result in results] == [0.7, 0.5]
    assert [result.retrieval_rank for result in results] == [1, 2]
    assert results[0].dense_rank is None
    query = instance.client.calls[0]
    assert query["with_payload"] is True
    assert query["with_vectors"] is False
    assert {condition.key for condition in query["query_filter"].must} == {"user_id", "chat_id", "document_id"}
    assert all(prefetch.filter == query["query_filter"] for prefetch in query["prefetch"])
    assert instance.search("user-1", "question", np.array([0.1, 0.2]), "chat-1") == "1. old\n\n2. new"


def test_qdrant_writes_canonical_chunk_ids_and_metadata() -> None:
    instance = qdrant()
    document = Document(document_id="doc-1", user_id="user-1", chat_id="chat-1", type=".pdf", file_name="notes.pdf")
    chunk = Chunk(chunk_id=str(uuid4()), document_id="doc-1", content="evidence", page=3, chunk_index=4, ocr_used=True)
    instance.add(np.array([[0.1, 0.2]], dtype=np.float32), [chunk.content], "user-1", document, "chat-1", [chunk])
    batch = instance.client.calls[0]["points"]
    assert batch.ids == [chunk.chunk_id]
    assert batch.payloads[0]["page"] == 3
    assert batch.payloads[0]["chunk_id"] == chunk.chunk_id
    assert batch.payloads[0]["file_name"] == "notes.pdf"
    with pytest.raises(ValueError, match="counts"):
        instance.add(np.array([]), [chunk.content], "user-1", document, "chat-1", [chunk])
    with pytest.raises(ValueError, match="unique"):
        instance.add(np.array([[0.1], [0.2]]), [chunk.content] * 2, "user-1", document, "chat-1", [chunk, chunk])


class MemorySql:
    def __init__(self):
        self.updated = None
        self.chunks = []
        self.documents = []

    def select_chat_history(self, **kwargs):
        return {"summary": "Previous summary", "conversation": []}

    def select_document_by_chat(self, user_id, chat_id):
        assert (user_id, chat_id) == ("user-1", "chat-1")
        return [{"document_id": "doc-1", "file_name": "legacy.pdf"}]

    def update_chat_history(self, **kwargs):
        self.updated = kwargs

    def insert_document(self, doc):
        self.documents.append(doc)

    def insert_chunks(self, chunks):
        self.chunks.extend(chunks)


class EvidenceChatbot:
    def stream(self, **kwargs):
        assert "SOURCE 1" in kwargs["data"]
        assert "Page: 3" in kwargs["data"]
        assert "legacy.pdf" in kwargs["data"]
        yield "Fact [[sour"
        yield "ce:1]] and [[source:99]]"

    def summarize_conversation(self, **kwargs):
        return "Updated summary"


def source_pipeline() -> Pipeline:
    return Pipeline(
        sql=MemorySql(),
        qdrant=SimpleNamespace(search_hits=lambda **kwargs: [hit(file_name="")]),
        embedding_model=SimpleNamespace(embed_query=lambda _: [0.1, 0.2]),
        chatbot=EvidenceChatbot(),
    )


def test_pipeline_snapshots_the_same_evidence_and_validates_after_save() -> None:
    pipeline = source_pipeline()
    events = list(pipeline.query_events("user-1", "question", "chat-1", include_trace=True))
    assert [event.event for event in events] == ["sources", "retrieval_trace", "token", "token", "citations"]
    metadata = pipeline.sql.updated["source_metadata"]
    assert metadata["sources"] == events[0].payload["sources"]
    assert metadata["cited_source_ids"] == [1]
    assert metadata["invalid_source_ids"] == [99]
    assert events[-1].payload["saved"] is True
    assert "content" not in events[1].payload["hits"][0]


def test_failed_history_save_never_emits_citations_or_done() -> None:
    pipeline = source_pipeline()

    def fail(**kwargs):
        raise RuntimeError("private storage failure")

    pipeline.sql.update_chat_history = fail
    events = list(stream_retrieval_events(pipeline, "user-1", "question", "chat-1", include_sources=True))
    assert any("event: sources" in event for event in events)
    assert "event: error" in events[-1]
    assert not any("event: done" in event or "event: citations" in event for event in events)
    assert "private storage failure" not in "".join(events)


def test_ingestion_preserves_pages_across_batch_boundary(monkeypatch) -> None:
    pages = [{"page": 3, "texts": ["chunk"] * 200, "ocr_used": True}, {"page": 4, "texts": ["last"]}]
    monkeypatch.setattr(extractor.ExtractorFactory, "create", staticmethod(lambda _: SimpleNamespace(extract=lambda _: pages)))
    batches = []
    pipeline = Pipeline(
        sql=MemorySql(), qdrant=SimpleNamespace(add=lambda **kwargs: batches.append(kwargs)),
        embedding_model=SimpleNamespace(embed_passages=lambda texts: np.zeros((len(texts), 2), dtype=np.float32)),
    )
    pipeline.insert_doc_pipeline("notes.pdf", "user-1", "chat-1", "notes.pdf")
    assert [len(batch["chunks"]) for batch in batches] == [200, 1]
    assert [chunk.chunk_index for chunk in pipeline.sql.chunks] == list(range(201))
    assert pipeline.sql.chunks[199].page == 3
    assert pipeline.sql.chunks[200].page == 4
    assert batches[0]["chunks"][0].chunk_id == pipeline.sql.chunks[0].chunk_id


def test_txt_docx_extractors_share_pipeline_contract(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(extractor, "chunking", lambda text: [text])
    text_file = tmp_path / "notes.txt"
    text_file.write_text("Evidence", encoding="utf-8")
    assert extractor.TextExtractor().extract(text_file) == [{"page": None, "texts": ["Evidence"]}]
    from docx import Document as WordDocument
    word_file = tmp_path / "notes.docx"
    document = WordDocument()
    document.add_paragraph("Evidence")
    document.save(word_file)
    assert extractor.WordExtractor().extract(word_file) == [{"page": None, "texts": ["Evidence"]}]


def test_agent_sources_are_available_before_final_text_and_narration_is_hidden() -> None:
    class StubAgent:
        def stream(self, _, context, **kwargs):
            yield {"type": "messages", "data": (AIMessageChunk(content="Private narration"), {"langgraph_node": "model"})}
            yield {"type": "updates", "data": {"model": {"messages": [AIMessage(content="Private narration", tool_calls=[{"id": "tool-1", "name": "qdrant_query", "args": {}}])]}}}
            context.sources.register([hit()])
            yield {"type": "updates", "data": {"tools": {}}}
            yield {"type": "messages", "data": (AIMessageChunk(content="Fact [[source:1]]"), {"langgraph_node": "model"})}
            yield {"type": "updates", "data": {"model": {"messages": [AIMessage(content="Fact [[source:1]]")]}}}

    agent = Agents.__new__(Agents)
    agent._agent = StubAgent()
    events = list(agent.stream_events("question", "", "user-1", "chat-1", SourceRegistry(), {}))
    assert [event.event for event in events] == ["sources", "token"]
    assert events[-1].payload["content"] == "Fact [[source:1]]"


def test_document_tool_registers_sources_in_runtime_only() -> None:
    tools = ToolList.__new__(ToolList)
    tools.embedding_model = SimpleNamespace(embed_query=lambda _: [0.1])
    tools.qdrant = SimpleNamespace(search_hits=lambda **kwargs: [hit(file_name="")])
    runtime = SimpleNamespace(context=AppContext("user-1", "chat-1", SourceRegistry(), {"doc-1": "legacy.pdf"}))
    result = tools._qdrant_query_tool("question", runtime)
    assert "SOURCE 1" in result and "legacy.pdf" in result
    assert len(runtime.context.sources.snapshot()) == 1


def test_sql_saves_source_snapshot_without_rewriting_legacy_turns() -> None:
    class Query:
        def __init__(self):
            self.updated = None
            self.inserted = None

        def select(self, _): return self
        def eq(self, *args): return self
        def single(self): return self
        def execute(self): return SimpleNamespace(data={"conversation": [{"user": "old", "chatbot": "old answer"}]})
        def update(self, data): self.updated = data; return self
        def insert(self, data): self.inserted = data; return self

    query = Query()
    sql = Supabase_Manager.__new__(Supabase_Manager)
    sql.supabase = SimpleNamespace(table=lambda _: query)
    metadata = {"turn_id": "turn-1", "sources": SourceRegistry().register([hit()])[0].model_dump()}
    sql.update_chat_history("user-1", "chat-1", "question", "answer", "summary", metadata)
    assert query.updated["conversation"][0] == {"user": "old", "chatbot": "old answer"}
    assert query.updated["conversation"][1]["turn_id"] == "turn-1"
    sql.insert_chunks([Chunk(chunk_id="chunk-1", document_id="doc-1", content="source", page=3, chunk_index=4)])
    assert query.inserted[0]["page"] == 3
    assert query.inserted[0]["chunk_index"] == 4


def test_local_qdrant_native_hybrid_preserves_metadata_and_tenant_scope() -> None:
    from qdrant_client import QdrantClient
    from qdrant_client.models import Distance, Modifier, SparseVectorParams, VectorParams

    instance = qdrant()
    instance.client = QdrantClient(":memory:")
    instance.client.create_collection(
        "user_documents", vectors_config={"dense": VectorParams(size=2, distance=Distance.COSINE)},
        sparse_vectors_config={"bm25": SparseVectorParams(modifier=Modifier.IDF)},
    )
    for user, chat in [("user-1", "chat-1"), ("user-2", "chat-1"), ("user-1", "chat-2")]:
        document = Document(document_id=str(uuid4()), user_id=user, chat_id=chat, type=".pdf", file_name="scoped.pdf")
        chunk = Chunk(chunk_id=str(uuid4()), document_id=document.document_id, content=f"Evidence for {user}/{chat}", page=5, chunk_index=0)
        instance.add(np.array([[1.0, 0.1]], dtype=np.float32), [chunk.content], user, document, chat, [chunk])
    results = instance.search_hits("user-1", "question", np.array([1.0, 0.1]), "chat-1")
    assert len(results) == 1
    assert results[0].content == "Evidence for user-1/chat-1"
    assert results[0].chunk_id == results[0].qdrant_point_id
    assert results[0].page == 5
    assert results[0].score > 0
    assert results[0].score_type == "rrf"
    instance.client.close()


def test_real_langgraph_document_tool_uses_request_registry() -> None:
    from langchain.agents import create_agent
    from langchain_core.language_models.fake_chat_models import FakeMessagesListChatModel
    from langchain_core.tools import StructuredTool

    class ToolCallingModel(FakeMessagesListChatModel):
        def bind_tools(self, tools, **kwargs):
            return self

    tools = ToolList.__new__(ToolList)
    tools.embedding_model = SimpleNamespace(embed_query=lambda _: [0.1])
    tools.qdrant = SimpleNamespace(search_hits=lambda **kwargs: [hit()])
    model = ToolCallingModel(responses=[
        AIMessage(content="Pre-tool narration", tool_calls=[{
            "id": "tool-1", "name": "qdrant_query", "args": {"query": "question"},
        }]),
        AIMessage(content="Fact [[source:1]]"),
    ])
    agent = Agents.__new__(Agents)
    agent._agent = create_agent(
        model=model, tools=[StructuredTool.from_function(tools._qdrant_query_tool, name="qdrant_query")],
        context_schema=AppContext,
    )
    events = list(agent.stream_events("question", "", "user-1", "chat-1", SourceRegistry(), {}))
    assert events[0].event == "sources"
    assert events[0].payload["sources"][0]["page"] == 3
    assert "".join(event.payload["content"] for event in events if event.event == "token") == "Fact [[source:1]]"
