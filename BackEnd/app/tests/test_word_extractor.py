"""Integration and performance test for :class:`WordExtractor`."""

from __future__ import annotations

import json
import resource
import time
from pathlib import Path
from typing import Any

from docx import Document

from BackEnd.app.doc_extractor.extractor import (
    ExtractorFactory,
    WordExtractor,
)


SAMPLE_PARAGRAPHS = [
    "Word Extractor Integration Test",
    "Đây là đoạn văn tiếng Việt dùng để kiểm tra Unicode.",
    "The extractor should preserve paragraph order and join them with newlines.",
]
BYTES_PER_MIB = 1024 * 1024


def _current_rss_bytes() -> int:
    """Return this process's current resident set size on Linux."""
    status_path = Path("/proc/self/status")
    for line in status_path.read_text(encoding="utf-8").splitlines():
        if line.startswith("VmRSS:"):
            return int(line.split()[1]) * 1024

    raise RuntimeError("VmRSS was not found in /proc/self/status")


def _peak_rss_bytes() -> int:
    """Return peak resident memory for this process on Linux."""
    usage = resource.getrusage(resource.RUSAGE_SELF)
    return usage.ru_maxrss * 1024


def _mib(byte_count: int) -> float:
    return byte_count / BYTES_PER_MIB


def _create_sample_document(file_path: Path) -> None:
    """Create a representative Word document for extraction testing."""
    document = Document()
    document.add_heading(SAMPLE_PARAGRAPHS[0], level=1)
    for paragraph in SAMPLE_PARAGRAPHS[1:]:
        document.add_paragraph(paragraph)
    document.save(file_path)


def _print_benchmark_report(
    *,
    file_path: Path,
    output: list[dict[str, Any]],
    elapsed_seconds: float,
    rss_before_bytes: int,
    rss_after_bytes: int,
    peak_rss_bytes: int,
) -> None:
    """Print benchmark metrics and all extracted Word output records."""
    report = {
        "file_path": str(file_path),
        "output_count": len(output),
        "elapsed_seconds": round(elapsed_seconds, 6),
        "rss_before_mib": round(_mib(rss_before_bytes), 3),
        "rss_after_mib": round(_mib(rss_after_bytes), 3),
        "rss_delta_mib": round(_mib(rss_after_bytes - rss_before_bytes), 3),
        "peak_rss_mib": round(_mib(peak_rss_bytes), 3),
    }

    print("\nWordExtractor benchmark")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    print("\nWordExtractor output")
    print(json.dumps(output, ensure_ascii=False, indent=2))


def test_extract_word_document(tmp_path: Path) -> None:
    """Extract a real DOCX fixture and report runtime, memory, and output."""
    file_path = tmp_path / "word_extractor_sample.docx"
    _create_sample_document(file_path)

    extractor = ExtractorFactory.create(str(file_path))
    assert isinstance(extractor, WordExtractor)

    rss_before_bytes = _current_rss_bytes()
    started_at = time.perf_counter()
    output = extractor.extract(str(file_path))
    elapsed_seconds = time.perf_counter() - started_at
    rss_after_bytes = _current_rss_bytes()
    peak_rss_bytes = max(
        _peak_rss_bytes(),
        rss_before_bytes,
        rss_after_bytes,
    )

    assert output == [
        {
            "page": None,
            "text": "\n".join(SAMPLE_PARAGRAPHS),
        }
    ]

    _print_benchmark_report(
        file_path=file_path,
        output=output,
        elapsed_seconds=elapsed_seconds,
        rss_before_bytes=rss_before_bytes,
        rss_after_bytes=rss_after_bytes,
        peak_rss_bytes=peak_rss_bytes,
    )
